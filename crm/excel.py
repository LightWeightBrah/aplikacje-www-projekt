"""Import i eksport Excela (.xlsx).

- import: manager wgrywa plik z bazą kontaktów, powstaje paczka (LeadBatch)
  z leadami w puli; duplikaty (po numerze telefonu) i błędne wiersze są
  pomijane i opisane w raporcie,
- eksport: lista leadów do pliku .xlsx (z filtrami z API).
"""

import io
import json
import re
import unicodedata
from datetime import datetime
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.renderers import BaseRenderer

from .models import Company, Lead, LeadBatch

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_IMPORT_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_IMPORT_ROWS = 5000

# Nazwy kolumn, które rozpoznajemy (po usunięciu polskich znaków, spacji i znaków specjalnych)
HEADER_ALIASES = {
    "contact_name": (
        "imieinazwisko", "imienazwisko", "kontakt", "osoba", "osobakontaktowa",
        "contactname", "fullname",
    ),
    "phone": (
        "telefon", "tel", "numertelefonu", "nrtelefonu", "nrtel", "komorka",
        "phone", "phonenumber",
    ),
    "email": ("email", "mail", "emailadres", "adresemail"),
    "company": ("firma", "nazwafirmy", "company", "companyname"),
    "nip": ("nip", "numernip"),
    "krs": ("krs", "numerkrs"),
    "regon": ("regon",),
    "city": ("miasto", "city", "miejscowosc"),
}
REQUIRED_COLUMNS = {"contact_name": "imię i nazwisko", "phone": "telefon"}


class ImportFileError(Exception):
    """Plik nie nadaje się do importu (zły format, brak kolumn, za duży)."""


class RowError(Exception):
    """Jeden wiersz jest błędny (reszta pliku jest importowana dalej)."""


# ------------------------------------------------------------------ pomocnicze

def _header_key(value):
    text = str(value or "").lower().replace("ł", "l")
    text = unicodedata.normalize("NFKD", text)
    return re.sub(r"[^a-z0-9]", "", text)


def _text(value):
    """Wartość komórki jako tekst (liczby z Excela bez końcówki .0)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _id_digits(value, length):
    """NIP/KRS/REGON: same cyfry; liczbie z Excela odzyskujemy zera z przodu."""
    digits = re.sub(r"\D", "", _text(value))
    if digits and isinstance(value, (int, float)) and len(digits) < length:
        digits = digits.zfill(length)
    return digits


def phone_key(phone):
    """Postać numeru do porównywania duplikatów: same cyfry, bez prefiksu +48 / 0048."""
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("0048"):
        digits = digits[4:]
    elif len(digits) == 11 and digits.startswith("48"):
        digits = digits[2:]
    return digits


def _validate_ids(record):
    for field_name in ("nip", "krs", "regon"):
        value = record[field_name]
        if not value:
            continue
        try:
            Company._meta.get_field(field_name).run_validators(value)
        except ValidationError as error:
            raise RowError(error.messages[0])


def _clean_row(raw):
    """Sprawdza jeden wiersz i zwraca oczyszczone dane albo zgłasza RowError."""
    record = {
        "contact_name": _text(raw.get("contact_name")),
        "phone": _text(raw.get("phone")),
        "email": _text(raw.get("email")),
        "company": _text(raw.get("company")),
        "city": _text(raw.get("city")),
        "nip": _id_digits(raw.get("nip"), 10),
        "krs": _id_digits(raw.get("krs"), 10),
        "regon": _id_digits(raw.get("regon"), 9),
    }
    if not record["contact_name"]:
        raise RowError("Brak imienia i nazwiska.")
    if len(record["contact_name"]) > 150:
        raise RowError("Imię i nazwisko jest za długie (max 150 znaków).")
    if not 9 <= len(re.sub(r"\D", "", record["phone"])) <= 15:
        raise RowError("Numer telefonu musi mieć od 9 do 15 cyfr.")
    if len(record["phone"]) > 30:
        raise RowError("Numer telefonu jest za długi (max 30 znaków).")
    if record["email"]:
        try:
            validate_email(record["email"])
        except ValidationError:
            raise RowError("Niepoprawny adres e-mail.")
    if len(record["company"]) > 200:
        raise RowError("Nazwa firmy jest za długa (max 200 znaków).")
    if len(record["city"]) > 100:
        raise RowError("Nazwa miasta jest za długa (max 100 znaków).")
    _validate_ids(record)
    return record


def _find_or_create_company(record, cache):
    """Firma po NIP (a bez NIP po nazwie); nowa powstaje tylko, gdy jej nie ma."""
    nip, name = record["nip"], record["company"]
    if not (nip or name):
        return None

    key = ("nip", nip) if nip else ("name", name.lower())
    if key in cache:
        return cache[key]

    if nip:
        company = Company.objects.filter(nip=nip).first()
    else:
        company = Company.objects.filter(name__iexact=name, nip="").first()
    if company is None:
        company = Company.objects.create(
            name=name or f"Firma NIP {nip}",
            nip=nip,
            krs=record["krs"],
            regon=record["regon"],
            city=record["city"],
        )
    cache[key] = company
    return company


def _read_sheet(uploaded_file):
    """Otwiera plik i zwraca (mapa kolumn, iterator wierszy)."""
    if uploaded_file.size > MAX_IMPORT_BYTES:
        raise ImportFileError("Plik jest za duży (max 5 MB).")
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        rows = workbook.worksheets[0].iter_rows(values_only=True)
    except (BadZipFile, InvalidFileException, KeyError, IndexError, OSError, ValueError):
        raise ImportFileError("Nie udało się odczytać pliku. Wgraj poprawny plik .xlsx.")

    header = next(rows, None)
    if not header:
        raise ImportFileError("Plik jest pusty.")

    columns = {}
    for index, title in enumerate(header):
        key = _header_key(title)
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases and field_name not in columns:
                columns[field_name] = index
    missing = [label for name, label in REQUIRED_COLUMNS.items() if name not in columns]
    if missing:
        raise ImportFileError(
            "Brak wymaganych kolumn: " + ", ".join(missing) + ". "
            "Pierwszy wiersz pliku musi zawierać nagłówki."
        )
    return columns, rows


# ---------------------------------------------------------------------- import

def import_leads(uploaded_file, user, batch_name="", product=None):
    """Importuje leady z pliku .xlsx do nowej paczki (LeadBatch).

    Zwraca raport: liczba utworzonych leadów, pominięte duplikaty i błędne wiersze.
    Gdy nie powstał żaden lead, nic nie jest zapisywane (batch = None).
    """
    if not batch_name:
        stem = re.sub(r"\.xlsx$", "", uploaded_file.name, flags=re.IGNORECASE)
        batch_name = f"Import {stem} {timezone.localtime():%Y-%m-%d %H:%M}"
    batch_name = batch_name[:150]
    if LeadBatch.objects.filter(name=batch_name).exists():
        raise ImportFileError(f"Baza o nazwie „{batch_name}” już istnieje.")

    columns, rows = _read_sheet(uploaded_file)

    # numery już obecne w systemie (jakikolwiek lead, także przypisany lub zamknięty)
    existing = {}
    for lead_id, phone in Lead.objects.values_list("id", "phone"):
        existing.setdefault(phone_key(phone), lead_id)
    seen_in_file = {}

    leads, duplicates, errors = [], [], []
    total_rows = 0
    company_cache = {}

    with transaction.atomic():
        batch = LeadBatch.objects.create(
            name=batch_name,
            source_file_name=uploaded_file.name[:255],
            created_by=user,
        )
        for number, row in enumerate(rows, start=2):
            if all(_text(cell) == "" for cell in row):
                continue  # pusty wiersz
            total_rows += 1
            if total_rows > MAX_IMPORT_ROWS:
                raise ImportFileError(f"Za dużo wierszy (max {MAX_IMPORT_ROWS}).")

            raw = {
                name: (row[index] if index < len(row) else None)
                for name, index in columns.items()
            }
            try:
                record = _clean_row(raw)
            except RowError as error:
                errors.append({"row": number, "message": str(error)})
                continue

            key = phone_key(record["phone"])
            if key in existing:
                duplicates.append({
                    "row": number,
                    "phone": record["phone"],
                    "reason": f"Numer już istnieje w systemie (lead #{existing[key]}).",
                })
                continue
            if key in seen_in_file:
                duplicates.append({
                    "row": number,
                    "phone": record["phone"],
                    "reason": f"Numer powtórzony w pliku (wiersz {seen_in_file[key]}).",
                })
                continue
            seen_in_file[key] = number

            leads.append(
                Lead(
                    batch=batch,
                    company=_find_or_create_company(record, company_cache),
                    product=product,
                    contact_name=record["contact_name"],
                    phone=record["phone"],
                    email=record["email"],
                )
            )

        if leads:
            Lead.objects.bulk_create(leads)
        else:
            transaction.set_rollback(True)  # nic do zapisania: bez pustej paczki

    return {
        "batch": {"id": batch.pk, "name": batch.name} if leads else None,
        "total_rows": total_rows,
        "created": len(leads),
        "duplicates": duplicates,
        "errors": errors,
    }


# --------------------------------------------------------------------- eksport

def _local(value):
    """Data do Excela: lokalny czas bez strefy (Excel nie zna stref czasowych)."""
    if value is None:
        return None
    return timezone.localtime(value).replace(tzinfo=None)


EXPORT_COLUMNS = (
    ("ID", 8, lambda l: l.pk),
    ("Imię i nazwisko", 24, lambda l: l.contact_name),
    ("Telefon", 18, lambda l: l.phone),
    ("E-mail", 28, lambda l: l.email),
    ("Firma", 28, lambda l: l.company.name if l.company else ""),
    ("NIP", 13, lambda l: l.company.nip if l.company else ""),
    ("KRS", 13, lambda l: l.company.krs if l.company else ""),
    ("Produkt", 22, lambda l: l.product.name if l.product else ""),
    ("Baza", 22, lambda l: l.batch.name if l.batch else ""),
    ("Status", 18, lambda l: l.get_status_display()),
    ("Opiekun", 14, lambda l: l.owner.username if l.owner else "(pula)"),
    ("Próby", 8, lambda l: l.attempts_count),
    ("Następny kontakt", 18, lambda l: _local(l.next_contact_at)),
    ("Notatka kontaktu", 30, lambda l: l.next_contact_note),
    ("Ostatnia aktywność", 18, lambda l: _local(l.last_activity_at)),
)


def leads_to_xlsx(leads):
    """Buduje plik .xlsx z listą leadów i zwraca go jako bajty.

    Teksty zapisujemy zawsze jako tekst (nigdy jako formuły), żeby dane
    zaczynające się od \"=\" nie były wykonywane przez Excela.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leady"

    for column, (title, width, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"

    for row, lead in enumerate(leads, start=2):
        for column, (_, _, getter) in enumerate(EXPORT_COLUMNS, start=1):
            value = getter(lead)
            cell = sheet.cell(row=row, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            elif isinstance(value, datetime):
                cell.number_format = "YYYY-MM-DD HH:MM"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class XlsxRenderer(BaseRenderer):
    """Odpowiedź jako plik Excela; ewentualne błędy zwracamy jako JSON."""

    media_type = XLSX_CONTENT_TYPE
    format = "xlsx"
    charset = None
    render_style = "binary"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        if isinstance(data, bytes):
            return data
        return json.dumps(data).encode("utf-8")