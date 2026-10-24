import io
from datetime import datetime, time, timedelta

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient, APITestCase

from .models import (
    Activity,
    AgentProfile,
    Company,
    Lead,
    LeadAssignment,
    LeadBatch,
    Product,
)
from .roles import AGENT, MANAGER, can_view_team_stats
from .services import release_stale_leads


# Szybszy hasher haseł, używany tylko w testach
@override_settings(
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"]
)
class CrmTestBase(APITestCase):
    """Wspólne dane testowe i funkcje pomocnicze."""

    def setUp(self):
        agent_group, _ = Group.objects.get_or_create(name=AGENT)
        manager_group, _ = Group.objects.get_or_create(name=MANAGER)

        self.agent1 = User.objects.create_user("agent1", password="haslo12345")
        self.agent2 = User.objects.create_user("agent2", password="haslo12345")
        self.manager = User.objects.create_user("manager1", password="haslo12345")
        self.agent1.groups.add(agent_group)
        self.agent2.groups.add(agent_group)
        self.manager.groups.add(manager_group)

        self.no_role = User.objects.create_user("bezroli", password="haslo12345")

    def api(self, user=None):
        client = APIClient()
        if user is not None:
            client.force_authenticate(user=user)
        return client

    def make_lead(self, owner=None, **kwargs):
        data = {"contact_name": "Jan Kowalski", "phone": "500100200"}
        data.update(kwargs)
        return Lead.objects.create(owner=owner, **data)


class AuthAndPermissionTests(CrmTestBase):
    """Uwierzytelnianie JWT, role i uprawnienia."""

    # Brak tokena daje 401
    def test_unauthenticated_request_returns_401(self):
        response = self.api().get("/api/leads/")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    # Token JWT z /api/token/ daje dostęp do API
    def test_jwt_token_allows_access_to_api(self):
        response = self.client.post(
            "/api/token/",
            {"username": "agent1", "password": "haslo12345"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("access", response.data)

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
        response = self.client.get("/api/my-leads/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    # Użytkownik bez roli Agent/Manager dostaje 403
    def test_user_without_role_gets_403(self):
        response = self.api(self.no_role).get("/api/leads/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    # Produkty może dodawać tylko manager
    def test_agent_cannot_create_product_but_manager_can(self):
        payload = {"name": "Światłowód 1 Gb"}

        response = self.api(self.agent1).post("/api/products/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

        response = self.api(self.manager).post("/api/products/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    # Leada może usunąć tylko manager
    def test_only_manager_can_delete_lead(self):
        lead = self.make_lead(owner=self.agent1)

        response = self.api(self.agent1).delete(f"/api/leads/{lead.id}/")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(Lead.objects.filter(id=lead.id).exists())

        response = self.api(self.manager).delete(f"/api/leads/{lead.id}/")
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Lead.objects.filter(id=lead.id).exists())


class LeadVisibilityTests(CrmTestBase):
    """Widoczność leadów zależnie od roli."""

    def setUp(self):
        super().setUp()
        self.lead_a1 = self.make_lead(owner=self.agent1, contact_name="Lead agenta 1")
        self.lead_a2 = self.make_lead(owner=self.agent2, contact_name="Lead agenta 2")
        self.lead_pool = self.make_lead(owner=None, contact_name="Lead z puli")

    # Agent widzi tylko swoje leady
    def test_agent_sees_only_own_leads(self):
        response = self.api(self.agent1).get("/api/leads/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        ids = [item["id"] for item in response.data["results"]]
        self.assertEqual(ids, [self.lead_a1.id])

    # Manager widzi wszystkie leady
    def test_manager_sees_all_leads(self):
        response = self.api(self.manager).get("/api/leads/")
        ids = {item["id"] for item in response.data["results"]}
        self.assertEqual(
            ids, {self.lead_a1.id, self.lead_a2.id, self.lead_pool.id}
        )

    # Agent nie otworzy cudzego leada (404)
    def test_agent_cannot_open_someone_elses_lead(self):
        response = self.api(self.agent1).get(f"/api/leads/{self.lead_a2.id}/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    # Lead agenta należy do niego, lead managera trafia do puli
    def test_agent_creates_lead_as_owner_manager_creates_into_pool(self):
        payload = {"contact_name": "Nowy Klient", "phone": "600700800"}

        response = self.api(self.agent1).post("/api/leads/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["owner"], self.agent1.id)

        response = self.api(self.manager).post("/api/leads/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIsNone(response.data["owner"])


class LeadValidationAndSearchTests(CrmTestBase):
    """Walidacja danych i wyszukiwanie leadów."""

    # Za krótki numer telefonu jest odrzucany
    def test_phone_too_short_is_rejected(self):
        response = self.api(self.agent1).post(
            "/api/leads/",
            {"contact_name": "Ktoś", "phone": "123"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("phone", response.data)

    # Status callback wymaga terminu, a sprawa (notatka) jest opcjonalna
    def test_callback_status_requires_date_but_not_note(self):
        response = self.api(self.agent1).post(
            "/api/leads/",
            {"contact_name": "Ktoś", "phone": "500100200", "status": "callback"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        response = self.api(self.agent1).post(
            "/api/leads/",
            {
                "contact_name": "Ktoś",
                "phone": "500100200",
                "status": "callback",
                "next_contact_at": (timezone.now() + timedelta(days=1)).isoformat(),
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    # Wyszukiwanie po NIP i KRS firmy
    def test_search_finds_lead_by_company_nip_and_krs(self):
        company = Company.objects.create(
            name="Test Sp. z o.o.", nip="1234567890", krs="0000123456"
        )
        wanted = self.make_lead(owner=self.agent1, company=company)
        self.make_lead(owner=self.agent1, contact_name="Inny Lead")
        client = self.api(self.agent1)

        for query in ("1234567890", "0000123456"):
            response = client.get("/api/leads/", {"search": query})
            ids = [item["id"] for item in response.data["results"]]
            self.assertEqual(ids, [wanted.id], msg=f"search={query}")


class NextLeadTests(CrmTestBase):
    """Endpoint leads/next."""

    # Lead z puli jest przypisywany agentowi
    def test_next_assigns_lead_from_pool_to_agent(self):
        lead = self.make_lead(owner=None)

        response = self.api(self.agent1).post("/api/leads/next/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["id"], lead.id)
        lead.refresh_from_db()
        self.assertEqual(lead.owner, self.agent1)
        self.assertIsNotNone(lead.claimed_at)
        self.assertTrue(
            LeadAssignment.objects.filter(lead=lead, to_agent=self.agent1).exists()
        )

    # Pusta pula daje 404
    def test_next_returns_404_when_pool_is_empty(self):
        self.make_lead(owner=self.agent2)

        response = self.api(self.agent1).post("/api/leads/next/")

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    # Zaległe oddzwonienie blokuje losowanie (409)
    def test_next_blocked_by_overdue_callback_returns_409(self):
        self.make_lead(
            owner=self.agent1,
            status=Lead.Status.CALLBACK,
            next_contact_at=timezone.now() - timedelta(hours=1),
            next_contact_note="Oddzwonić w sprawie oferty",
        )
        pool_lead = self.make_lead(owner=None, contact_name="Lead z puli")

        response = self.api(self.agent1).post("/api/leads/next/")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.data["overdue_count"], 1)
        pool_lead.refresh_from_db()
        self.assertIsNone(pool_lead.owner)

    # Agent nie dostaje leada, który już prowadził
    def test_next_does_not_return_lead_agent_already_handled(self):
        lead = self.make_lead(owner=None)
        LeadAssignment.objects.create(
            lead=lead, from_agent=self.agent1, reason="Zwrot do puli"
        )

        response = self.api(self.agent1).post("/api/leads/next/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

        response = self.api(self.agent2).post("/api/leads/next/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["id"], lead.id)


class ActivityEffectTests(CrmTestBase):
    """Wpływ aktywności na lead."""

    def setUp(self):
        super().setUp()
        self.lead = self.make_lead(owner=self.agent1)

    def add_activity(self, user=None, **payload):
        payload.setdefault("lead", self.lead.id)
        return self.api(user or self.agent1).post(
            "/api/activities/", payload, format="json"
        )

    # Brak odpowiedzi zwiększa próby i ustawia ostatnią aktywność
    def test_no_answer_increases_attempts_and_sets_last_activity(self):
        response = self.add_activity(outcome="no_answer")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.attempts_count, 1)
        self.assertIsNotNone(self.lead.last_activity_at)
        self.assertEqual(self.lead.status, Lead.Status.NEW)
        self.assertEqual(Activity.objects.filter(lead=self.lead).count(), 1)

    # Notatka nie liczy się jako próba
    def test_note_does_not_count_as_attempt(self):
        self.add_activity(outcome="note", note="Klient prosił o maila")

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.attempts_count, 0)
        self.assertIsNotNone(self.lead.last_activity_at)

    # Prośba o oddzwonienie ustawia status i termin
    def test_callback_sets_status_and_next_contact(self):
        when = timezone.now() + timedelta(days=1)

        response = self.add_activity(
            outcome="callback",
            next_contact_at=when.isoformat(),
            next_contact_note="Oddzwonić jutro",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.status, Lead.Status.CALLBACK)
        self.assertIsNotNone(self.lead.next_contact_at)
        self.assertEqual(self.lead.next_contact_note, "Oddzwonić jutro")

    # Oddzwonienie bez terminu jest odrzucane
    def test_callback_without_date_is_rejected(self):
        response = self.add_activity(outcome="callback")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.attempts_count, 0)
        self.assertEqual(Activity.objects.count(), 0)

    # Termin z przeszłości jest odrzucany
    def test_callback_date_in_the_past_is_rejected(self):
        past = timezone.now() - timedelta(hours=1)

        response = self.add_activity(
            outcome="callback",
            next_contact_at=past.isoformat(),
            next_contact_note="Za późno",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # Odmowa zamyka leada i czyści termin
    def test_refused_closes_lead_and_clears_next_contact(self):
        self.lead.next_contact_at = timezone.now() + timedelta(hours=2)
        self.lead.next_contact_note = "Stary termin"
        self.lead.save()

        self.add_activity(outcome="refused")

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.status, Lead.Status.LOST)
        self.assertIsNone(self.lead.next_contact_at)
        self.assertEqual(self.lead.next_contact_note, "")

    # Zły numer ustawia status nie dzwonić
    def test_wrong_number_sets_do_not_call(self):
        self.add_activity(outcome="wrong_number")

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.status, Lead.Status.DO_NOT_CALL)

    # Agent nie doda aktywności do cudzego leada
    def test_agent_cannot_add_activity_to_someone_elses_lead(self):
        response = self.add_activity(user=self.agent2, outcome="no_answer")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.attempts_count, 0)

    # Autorem aktywności jest zalogowany agent
    def test_activity_is_recorded_with_logged_in_agent(self):
        self.add_activity(outcome="no_answer", agent=self.agent2.id)

        activity = Activity.objects.get(lead=self.lead)
        self.assertEqual(activity.agent, self.agent1)


class SoldOutcomeTests(CrmTestBase):
    """Wynik rozmowy „Sprzedane”: lead jest wygrany i zamknięty."""

    def test_sold_closes_lead_as_won_and_clears_next_contact(self):
        lead = self.make_lead(
            owner=self.agent1,
            status=Lead.Status.INTERESTED,
            next_contact_at=timezone.now() + timedelta(days=1),
            next_contact_note="oferta",
        )
        response = self.api(self.agent1).post(
            "/api/activities/", {"lead": lead.id, "outcome": "sold"}, format="json"
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        lead.refresh_from_db()
        self.assertEqual(lead.status, Lead.Status.WON)
        self.assertIsNone(lead.next_contact_at)
        self.assertEqual(lead.next_contact_note, "")
        self.assertEqual(lead.attempts_count, 1)

    def test_sold_is_counted_in_stats_as_answered_call(self):
        lead = self.make_lead(owner=self.agent1)
        self.api(self.agent1).post(
            "/api/activities/", {"lead": lead.id, "outcome": "sold"}, format="json"
        )
        data = self.api(self.agent1).get("/api/stats/?days=1").data
        self.assertEqual(data["outcomes"]["sold"], 1)
        self.assertEqual(data["agents"][0]["sold"], 1)
        self.assertEqual(data["agents"][0]["answered"], 1)
        self.assertEqual(data["leads_by_status"]["won"], 1)


class MyLeadsDueTests(CrmTestBase):
    """Endpoint my-leads/due."""

    # Podział na zaległe i nadchodzące leady
    def test_due_splits_overdue_and_upcoming_leads(self):
        now = timezone.now()
        overdue = self.make_lead(
            owner=self.agent1, next_contact_at=now - timedelta(hours=1)
        )
        upcoming = self.make_lead(
            owner=self.agent1, next_contact_at=now + timedelta(minutes=5)
        )
        self.make_lead(owner=self.agent1, next_contact_at=now + timedelta(hours=3))
        self.make_lead(owner=self.agent2, next_contact_at=now - timedelta(hours=1))

        response = self.api(self.agent1).get("/api/my-leads/due/", {"minutes": 15})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["window_minutes"], 15)
        self.assertEqual([lead["id"] for lead in response.data["overdue"]], [overdue.id])
        self.assertEqual(
            [lead["id"] for lead in response.data["upcoming"]], [upcoming.id]
        )

    # Błędny parametr minutes daje 400
    def test_due_rejects_invalid_minutes(self):
        client = self.api(self.agent1)
        for value in ("abc", "-5", "5000"):
            response = client.get("/api/my-leads/due/", {"minutes": value})
            self.assertEqual(
                response.status_code, status.HTTP_400_BAD_REQUEST, msg=value
            )


class LeadLifecycleTests(CrmTestBase):
    """Cykl życia leada: limit prób, zwrot do puli, nieosiągalni, nietknięte leady."""

    def setUp(self):
        super().setUp()
        self.lead = self.make_lead(owner=self.agent1, claimed_at=timezone.now())

    def no_answer(self, agent=None):
        return self.api(agent or self.agent1).post(
            "/api/activities/",
            {"lead": self.lead.id, "outcome": "no_answer"},
            format="json",
        )

    # Poniżej limitu lead zostaje u agenta
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=3)
    def test_lead_stays_with_agent_below_attempt_limit(self):
        self.no_answer()
        self.no_answer()

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.owner, self.agent1)

    # Po wyczerpaniu limitu lead wraca do puli i zostaje wpis w historii przypisań
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=3)
    def test_lead_returns_to_pool_after_attempt_limit(self):
        for _ in range(3):
            self.no_answer()

        self.lead.refresh_from_db()
        self.assertIsNone(self.lead.owner)
        self.assertEqual(self.lead.status, Lead.Status.NEW)
        self.assertIsNone(self.lead.claimed_at)
        assignment = LeadAssignment.objects.get(lead=self.lead, to_agent=None)
        self.assertEqual(assignment.from_agent, self.agent1)

    # Agent, który wyczerpał limit, nie dostanie tego leada ponownie
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=2)
    def test_agent_does_not_get_released_lead_again(self):
        self.no_answer()
        self.no_answer()

        response = self.api(self.agent1).post("/api/leads/next/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

        response = self.api(self.agent2).post("/api/leads/next/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["id"], self.lead.id)

    # Lead z obietnicą oddzwonienia nie wraca do puli mimo nieodebranych telefonów
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=2)
    def test_callback_lead_is_not_released(self):
        Lead.objects.filter(pk=self.lead.pk).update(
            status=Lead.Status.CALLBACK,
            next_contact_at=timezone.now() + timedelta(days=1),
            next_contact_note="Oddzwonić",
        )
        self.no_answer()
        self.no_answer()

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.owner, self.agent1)
        self.assertEqual(self.lead.status, Lead.Status.CALLBACK)

    # Odmowa lub inny wynik niż brak odpowiedzi nie zwraca leada do puli
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=1)
    def test_other_outcomes_do_not_release_lead(self):
        self.api(self.agent1).post(
            "/api/activities/",
            {"lead": self.lead.id, "outcome": "note", "note": "Zadzwonić po 15"},
            format="json",
        )

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.owner, self.agent1)

    # Gdy tyle prób zrobiło już N różnych agentów, lead staje się nieosiągalny
    @override_settings(LEAD_MAX_ATTEMPTS_PER_AGENT=1, LEAD_UNREACHABLE_AFTER_AGENTS=2)
    def test_lead_becomes_unreachable_after_enough_agents(self):
        self.no_answer()  # agent1 wyczerpał limit, lead wraca do puli
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.status, Lead.Status.NEW)

        response = self.api(self.agent2).post("/api/leads/next/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.no_answer(agent=self.agent2)  # drugi agent bez skutku

        self.lead.refresh_from_db()
        self.assertEqual(self.lead.status, Lead.Status.UNREACHABLE)
        self.assertIsNone(self.lead.owner)

        # nieosiągalny lead nie trafia już do losowania
        self.assertEqual(
            self.api(self.manager).post("/api/leads/next/").status_code,
            status.HTTP_404_NOT_FOUND,
        )

    # Nietknięty od ponad limitu godzin lead wraca do puli
    @override_settings(LEAD_CLAIM_TIMEOUT_HOURS=24)
    def test_release_stale_returns_untouched_lead(self):
        Lead.objects.filter(pk=self.lead.pk).update(
            claimed_at=timezone.now() - timedelta(hours=30)
        )

        released = release_stale_leads()

        self.assertEqual(released, 1)
        self.lead.refresh_from_db()
        self.assertIsNone(self.lead.owner)
        self.assertTrue(
            LeadAssignment.objects.filter(
                lead=self.lead, from_agent=self.agent1, to_agent=None
            ).exists()
        )

    # Świeży lub ruszony przez agenta lead zostaje u agenta
    @override_settings(LEAD_CLAIM_TIMEOUT_HOURS=24)
    def test_release_stale_ignores_fresh_and_touched_leads(self):
        fresh = self.make_lead(owner=self.agent2, claimed_at=timezone.now())
        touched = self.make_lead(
            owner=self.agent2, claimed_at=timezone.now() - timedelta(hours=30)
        )
        Activity.objects.create(
            lead=touched, agent=self.agent2, outcome=Activity.Outcome.NOTE
        )
        Lead.objects.filter(pk=self.lead.pk).update(claimed_at=timezone.now())

        self.assertEqual(release_stale_leads(), 0)
        fresh.refresh_from_db()
        touched.refresh_from_db()
        self.assertEqual(fresh.owner, self.agent2)
        self.assertEqual(touched.owner, self.agent2)

    # Tryb --dry-run tylko liczy i niczego nie zmienia
    @override_settings(LEAD_CLAIM_TIMEOUT_HOURS=24)
    def test_release_stale_dry_run_changes_nothing(self):
        Lead.objects.filter(pk=self.lead.pk).update(
            claimed_at=timezone.now() - timedelta(hours=30)
        )

        self.assertEqual(release_stale_leads(dry_run=True), 1)
        self.lead.refresh_from_db()
        self.assertEqual(self.lead.owner, self.agent1)


class ManagerLeadManagementTests(CrmTestBase):
    """Przypisywanie leadów, przepisanie leadów agenta i pula nieosiągalnych."""

    def post_assign(self, lead, agent_id, user=None, **extra):
        return self.api(user or self.manager).post(
            f"/api/leads/{lead.id}/assign/",
            {"agent": agent_id, **extra},
            format="json",
        )

    # Agent nie może przypisywać leadów
    def test_agent_cannot_assign_lead(self):
        lead = self.make_lead()
        response = self.post_assign(lead, self.agent1.id, user=self.agent1)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    # Manager przypisuje lead z puli agentowi, zostaje wpis w historii
    def test_manager_assigns_pool_lead_to_agent(self):
        lead = self.make_lead()
        response = self.post_assign(lead, self.agent2.id, reason="Duży klient")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        lead.refresh_from_db()
        self.assertEqual(lead.owner, self.agent2)
        self.assertIsNotNone(lead.claimed_at)
        assignment = LeadAssignment.objects.get(lead=lead)
        self.assertEqual(assignment.to_agent, self.agent2)
        self.assertEqual(assignment.reason, "Duży klient")

    # Przepisanie zachowuje status i termin kontaktu
    def test_reassign_keeps_status_and_next_contact(self):
        when = timezone.now() + timedelta(days=1)
        lead = self.make_lead(
            owner=self.agent1,
            status=Lead.Status.CALLBACK,
            next_contact_at=when,
            next_contact_note="Oddzwonić",
        )
        self.post_assign(lead, self.agent2.id)

        lead.refresh_from_db()
        self.assertEqual(lead.owner, self.agent2)
        self.assertEqual(lead.status, Lead.Status.CALLBACK)
        self.assertEqual(lead.next_contact_at, when)
        self.assertEqual(
            LeadAssignment.objects.get(lead=lead).from_agent, self.agent1
        )

    # Nieosiągalny lead po przypisaniu agentowi znów jest "nowy"
    def test_assigning_unreachable_lead_makes_it_new(self):
        lead = self.make_lead(status=Lead.Status.UNREACHABLE)
        self.post_assign(lead, self.agent1.id)

        lead.refresh_from_db()
        self.assertEqual(lead.status, Lead.Status.NEW)

    # Zamkniętego leada nie można przepisać
    def test_closed_lead_cannot_be_assigned(self):
        lead = self.make_lead(owner=self.agent1, status=Lead.Status.LOST)
        response = self.post_assign(lead, self.agent2.id)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # Przypisać można tylko użytkownika z rolą
    def test_cannot_assign_user_without_role(self):
        lead = self.make_lead()
        response = self.post_assign(lead, self.no_role.id)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # Do puli wraca lead nowy; zainteresowany trzeba przepisać
    def test_return_to_pool_only_for_new_or_unreachable(self):
        new_lead = self.make_lead(owner=self.agent1)
        hot_lead = self.make_lead(owner=self.agent1, status=Lead.Status.INTERESTED)

        ok = self.post_assign(new_lead, None)
        bad = self.post_assign(hot_lead, None)

        self.assertEqual(ok.status_code, status.HTTP_200_OK)
        new_lead.refresh_from_db()
        self.assertIsNone(new_lead.owner)
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    # Masowe przepisanie: tylko aktywne leady odchodzącego agenta
    def test_bulk_reassign_moves_only_active_leads(self):
        active = self.make_lead(owner=self.agent1)
        closed = self.make_lead(owner=self.agent1, status=Lead.Status.LOST)
        other = self.make_lead(owner=self.agent2)

        response = self.api(self.manager).post(
            "/api/leads/reassign/",
            {"from_agent": self.agent1.id, "to_agent": self.agent2.id},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["reassigned"], 1)
        active.refresh_from_db()
        closed.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(active.owner, self.agent2)
        self.assertEqual(closed.owner, self.agent1)
        self.assertEqual(other.owner, self.agent2)

    # Masowe przepisanie wymaga dwóch różnych agentów i roli managera
    def test_bulk_reassign_validation_and_permission(self):
        same = self.api(self.manager).post(
            "/api/leads/reassign/",
            {"from_agent": self.agent1.id, "to_agent": self.agent1.id},
            format="json",
        )
        denied = self.api(self.agent1).post(
            "/api/leads/reassign/",
            {"from_agent": self.agent1.id, "to_agent": self.agent2.id},
            format="json",
        )
        self.assertEqual(same.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

    # Pulę nieosiągalnych widzi tylko manager; filtr ready pokazuje leady po karencji
    @override_settings(LEAD_UNREACHABLE_COOLDOWN_DAYS=30)
    def test_unreachable_pool_and_ready_filter(self):
        old = self.make_lead(
            status=Lead.Status.UNREACHABLE,
            last_activity_at=timezone.now() - timedelta(days=40),
        )
        self.make_lead(
            status=Lead.Status.UNREACHABLE,
            last_activity_at=timezone.now() - timedelta(days=5),
        )

        denied = self.api(self.agent1).get("/api/leads/unreachable/")
        all_ = self.api(self.manager).get("/api/leads/unreachable/")
        ready = self.api(self.manager).get("/api/leads/unreachable/?ready=true")

        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(all_.data["count"], 2)
        self.assertEqual(ready.data["count"], 1)
        self.assertEqual(ready.data["results"][0]["id"], old.id)

    # Zwrot do obiegu dotyczy tylko leadów po karencji
    @override_settings(LEAD_UNREACHABLE_COOLDOWN_DAYS=30)
    def test_recycle_unreachable_after_cooldown(self):
        old = self.make_lead(
            status=Lead.Status.UNREACHABLE,
            last_activity_at=timezone.now() - timedelta(days=40),
        )
        fresh = self.make_lead(
            status=Lead.Status.UNREACHABLE,
            last_activity_at=timezone.now() - timedelta(days=5),
        )

        denied = self.api(self.agent1).post("/api/leads/recycle-unreachable/")
        response = self.api(self.manager).post("/api/leads/recycle-unreachable/")

        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.data["recycled"], 1)
        old.refresh_from_db()
        fresh.refresh_from_db()
        self.assertEqual(old.status, Lead.Status.NEW)
        self.assertEqual(fresh.status, Lead.Status.UNREACHABLE)
        self.assertTrue(LeadAssignment.objects.filter(lead=old).exists())


class StatsTests(CrmTestBase):
    """Statystyki: lejek, skuteczność agentów, rozmowy dzień po dniu."""

    def setUp(self):
        super().setUp()
        self.make_lead()  # w puli
        self.make_lead(owner=self.agent1)
        self.make_lead(owner=self.agent1, status=Lead.Status.INTERESTED)
        self.make_lead(owner=self.agent2, status=Lead.Status.WON)
        lead = Lead.objects.filter(owner=self.agent1).first()
        for outcome in ("no_answer", "no_answer", "interested", "note"):
            Activity.objects.create(lead=lead, agent=self.agent1, outcome=outcome)

    def agent_row(self, response, username):
        return next(row for row in response.data["agents"] if row["username"] == username)

    # Manager dostaje lejek, pulę i wyniki wszystkich agentów
    def test_manager_sees_funnel_pool_and_all_agents(self):
        response = self.api(self.manager).get("/api/stats/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_leads"], 4)
        self.assertEqual(response.data["pool_leads"], 1)
        self.assertEqual(response.data["leads_by_status"]["new"], 2)
        self.assertEqual(response.data["leads_by_status"]["interested"], 1)
        self.assertEqual(response.data["leads_by_status"]["won"], 1)
        self.assertEqual(response.data["leads_by_status"]["unreachable"], 0)

        row = self.agent_row(response, "agent1")
        self.assertEqual(row["calls"], 3)  # notatka nie jest rozmową
        self.assertEqual(row["answered"], 1)
        self.assertEqual(row["interested"], 1)
        self.assertEqual(row["answer_rate"], 0.33)
        self.assertEqual(row["interest_rate"], 1.0)
        self.assertEqual(row["leads_total"], 2)
        self.assertEqual(self.agent_row(response, "agent2")["leads_won"], 1)
        self.assertIsNone(self.agent_row(response, "agent2")["answer_rate"])

    # Agent widzi tylko własne statystyki i nie widzi puli
    def test_agent_sees_only_own_stats(self):
        response = self.api(self.agent1).get("/api/stats/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["agents"]), 1)
        self.assertEqual(response.data["agents"][0]["username"], "agent1")
        self.assertIsNone(response.data["pool_leads"])
        self.assertEqual(response.data["total_leads"], 2)

    # Rozmowy dzień po dniu: długość = days, dzisiejsze rozmowy policzone
    def test_daily_calls_cover_each_day_in_window(self):
        response = self.api(self.manager).get("/api/stats/?days=3")

        daily = response.data["daily_calls"]
        self.assertEqual(len(daily), 3)
        self.assertEqual(daily[-1]["date"], timezone.localdate().isoformat())
        self.assertEqual(daily[-1]["calls"], 3)
        self.assertEqual(daily[0]["calls"], 0)

    # Okres jednodniowy dostaje rozbicie na godziny (24 pozycje), dłuższy – pustą listę
    def test_hourly_calls_only_for_single_day(self):
        today = self.api(self.manager).get("/api/stats/?days=1").data
        hours = today["hourly_calls"]
        self.assertEqual(len(hours), 24)
        self.assertEqual([h["hour"] for h in hours], list(range(24)))
        self.assertEqual(sum(h["calls"] for h in hours), 3)  # notatka nie jest rozmową

        week = self.api(self.manager).get("/api/stats/?days=7").data
        self.assertEqual(week["hourly_calls"], [])

    # Rozbicie godzinowe działa też dla jednodniowego zakresu dat i uwzględnia strefę czasową
    def test_hourly_calls_for_single_day_range_uses_local_time(self):
        today = timezone.localdate()
        ten_local = timezone.make_aware(datetime.combine(today, time(10, 15)))
        Activity.objects.update(created_at=ten_local)

        response = self.api(self.manager).get(f"/api/stats/?date_from={today}&date_to={today}")

        hours = response.data["hourly_calls"]
        self.assertEqual(hours[10]["calls"], 3)
        self.assertEqual(sum(h["calls"] for h in hours), 3)

    # Agent w rozbiciu godzinowym widzi tylko własne rozmowy
    def test_hourly_calls_for_agent_are_own_only(self):
        lead = Lead.objects.filter(owner=self.agent1).first()
        Activity.objects.create(lead=lead, agent=self.agent2, outcome="refused")

        mine = self.api(self.agent1).get("/api/stats/?days=1").data["hourly_calls"]

        self.assertEqual(sum(h["calls"] for h in mine), 3)

    # Rozmowy sprzed okna nie są liczone
    def test_old_activities_are_outside_window(self):
        Activity.objects.update(created_at=timezone.now() - timedelta(days=10))

        response = self.api(self.manager).get("/api/stats/?days=7")

        self.assertEqual(self.agent_row(response, "agent1")["calls"], 0)

    # Zły parametr days daje 400
    def test_invalid_days_returns_400(self):
        for value in ("0", "abc", "400"):
            response = self.api(self.manager).get(f"/api/stats/?days={value}")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # Rozbicie wyników rozmów w okresie (agent: tylko własne)
    def test_outcome_breakdown_for_period(self):
        lead = Lead.objects.filter(owner=self.agent1).first()
        for outcome in ("callback", "refused", "wrong_number"):
            Activity.objects.create(lead=lead, agent=self.agent1, outcome=outcome)
        Activity.objects.create(lead=lead, agent=self.agent2, outcome="refused")

        mine = self.api(self.agent1).get("/api/stats/?days=1").data
        self.assertEqual(mine["outcomes"]["no_answer"], 2)
        self.assertEqual(mine["outcomes"]["callback"], 1)
        self.assertEqual(mine["outcomes"]["interested"], 1)
        self.assertEqual(mine["outcomes"]["refused"], 1)  # odmowa agent2 nie jest moja
        self.assertEqual(mine["outcomes"]["wrong_number"], 1)
        self.assertEqual(mine["outcomes"]["note"], 1)

        team = self.api(self.manager).get("/api/stats/?days=1").data
        self.assertEqual(team["outcomes"]["refused"], 2)
        row = next(r for r in team["agents"] if r["username"] == "agent1")
        self.assertEqual(row["no_answer"], 2)
        self.assertEqual(row["callbacks"], 1)
        self.assertEqual(row["wrong_number"], 1)

    # Zaplanowane kontakty: zaległe, dziś, jutro, później
    def test_scheduled_counts(self):
        now = timezone.now()
        if timezone.localtime(now).hour == 23 and timezone.localtime(now).minute >= 50:
            self.skipTest("test zależy od pory dnia i nie jest wiarygodny tuż przed północą")
        local_today = timezone.localtime(now).replace(hour=23, minute=59)
        Lead.objects.filter(owner=self.agent1).update(next_contact_at=None)
        self.make_lead(owner=self.agent1, next_contact_at=now - timedelta(hours=2))
        self.make_lead(owner=self.agent1, next_contact_at=local_today)
        self.make_lead(owner=self.agent1, next_contact_at=now + timedelta(days=1))
        self.make_lead(owner=self.agent1, next_contact_at=now + timedelta(days=5))
        self.make_lead(owner=self.agent1, status=Lead.Status.LOST, next_contact_at=now + timedelta(days=5))

        scheduled = self.api(self.agent1).get("/api/stats/").data["scheduled"]
        self.assertEqual(scheduled["overdue"], 1)
        self.assertEqual(scheduled["later"], 1)
        self.assertEqual(sum(scheduled.values()), 4)  # przegrany lead się nie liczy

    # Zakres dat: tylko wybrane dni, włącznie z obiema datami
    def test_date_range_selects_days(self):
        today = timezone.localdate()
        yesterday = today - timedelta(days=1)
        old = Activity.objects.get(agent=self.agent1, outcome="interested")
        Activity.objects.filter(pk=old.pk).update(created_at=timezone.now() - timedelta(days=1))

        response = self.api(self.agent1).get(
            f"/api/stats/?date_from={yesterday.isoformat()}&date_to={yesterday.isoformat()}"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["days"], 1)
        self.assertEqual(response.data["date_from"], yesterday.isoformat())
        self.assertEqual(response.data["agents"][0]["calls"], 1)
        self.assertEqual(response.data["daily_calls"][0]["calls"], 1)

        today_only = self.api(self.agent1).get(
            f"/api/stats/?date_from={today.isoformat()}&date_to={today.isoformat()}"
        )
        self.assertEqual(today_only.data["agents"][0]["calls"], 2)

    # Zły zakres dat daje 400
    def test_invalid_date_range_returns_400(self):
        today = timezone.localdate()
        bad = (
            "date_from=2026-10-10",
            "date_from=abc&date_to=2026-10-10",
            f"date_from={today.isoformat()}&date_to={(today - timedelta(days=1)).isoformat()}",
            "date_from=2024-01-01&date_to=2026-01-01",
        )
        for query in bad:
            response = self.api(self.manager).get(f"/api/stats/?{query}")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, query)

    # Statystyki wymagają logowania i roli
    def test_stats_require_login_and_role(self):
        self.assertEqual(
            self.api().get("/api/stats/").status_code, status.HTTP_401_UNAUTHORIZED
        )
        self.assertEqual(
            self.api(self.no_role).get("/api/stats/").status_code,
            status.HTTP_403_FORBIDDEN,
        )


class ScheduleTests(CrmTestBase):
    """Plan dnia (my-leads/schedule) i kalendarz .ics."""

    def at(self, day, hour):
        return timezone.make_aware(datetime.combine(day, time(hour, 0)))

    def setUp(self):
        super().setUp()
        self.tomorrow = timezone.localdate() + timedelta(days=1)
        self.late = self.make_lead(
            owner=self.agent1,
            contact_name="Pozny Klient",
            status=Lead.Status.CALLBACK,
            next_contact_at=self.at(self.tomorrow, 14),
        )
        self.early = self.make_lead(
            owner=self.agent1,
            contact_name="Ranny Klient",
            status=Lead.Status.CALLBACK,
            next_contact_at=self.at(self.tomorrow, 9),
        )

    # Plan dnia: tylko ten dzień, posortowany po godzinie, tylko własne leady
    def test_schedule_lists_leads_of_the_day_in_order(self):
        self.make_lead(  # inny dzień
            owner=self.agent1,
            status=Lead.Status.CALLBACK,
            next_contact_at=self.at(self.tomorrow + timedelta(days=1), 10),
        )
        self.make_lead(  # cudzy lead
            owner=self.agent2,
            status=Lead.Status.CALLBACK,
            next_contact_at=self.at(self.tomorrow, 10),
        )
        self.make_lead(  # zamknięty lead
            owner=self.agent1,
            status=Lead.Status.LOST,
            next_contact_at=self.at(self.tomorrow, 11),
        )

        response = self.api(self.agent1).get(
            f"/api/my-leads/schedule/?date={self.tomorrow.isoformat()}"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        ids = [lead["id"] for lead in response.data["scheduled"]]
        self.assertEqual(ids, [self.early.id, self.late.id])
        self.assertEqual(response.data["overdue"], [])

    # Plan na dziś zawiera zaległe oddzwonienia z poprzednich dni
    def test_schedule_for_today_includes_overdue(self):
        overdue = self.make_lead(
            owner=self.agent1,
            status=Lead.Status.CALLBACK,
            next_contact_at=timezone.now() - timedelta(days=2),
        )

        response = self.api(self.agent1).get("/api/my-leads/schedule/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual([lead["id"] for lead in response.data["overdue"]], [overdue.id])
        self.assertEqual(response.data["date"], timezone.localdate().isoformat())

    # Zła data daje 400
    def test_schedule_invalid_date_returns_400(self):
        response = self.api(self.agent1).get("/api/my-leads/schedule/?date=jutro")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # Kalendarz .ics zawiera wydarzenia tylko własnych leadów
    def test_calendar_contains_only_own_future_callbacks(self):
        other = self.make_lead(
            owner=self.agent2,
            contact_name="Cudzy Klient",
            status=Lead.Status.CALLBACK,
            next_contact_at=self.at(self.tomorrow, 12),
        )

        response = self.api(self.agent1).get("/api/my-leads/calendar/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response["Content-Type"].startswith("text/calendar"))
        body = response.content.decode("utf-8")
        self.assertIn("BEGIN:VCALENDAR", body)
        self.assertEqual(body.count("BEGIN:VEVENT"), 2)
        self.assertIn("SUMMARY:Oddzwonić: Ranny Klient", body)
        self.assertNotIn(f"UID:lead-{other.id}@", body)

    # Plan dnia i kalendarz wymagają logowania
    def test_schedule_requires_login(self):
        self.assertEqual(
            self.api().get("/api/my-leads/schedule/").status_code,
            status.HTTP_401_UNAUTHORIZED,
        )


IMPORT_HEADERS = [
    "Imię i nazwisko", "Telefon", "E-mail", "Firma", "NIP", "KRS", "REGON", "Miasto",
]


def make_xlsx(rows, headers=IMPORT_HEADERS, filename="baza.xlsx"):
    """Buduje plik .xlsx w pamięci i zwraca go jako upload do testów."""
    workbook = Workbook()
    sheet = workbook.active
    if headers:
        sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        filename,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class ExcelImportTests(CrmTestBase):
    """Import bazy kontaktów z Excela."""

    def do_import(self, upload, user=None, **extra):
        return self.api(user or self.manager).post(
            "/api/batches/import/", {"file": upload, **extra}, format="multipart"
        )

    # Agent nie może importować
    def test_agent_cannot_import(self):
        response = self.do_import(make_xlsx([]), user=self.agent1)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    # Import tworzy paczkę, leady w puli i firmy (ta sama firma po NIP tylko raz)
    def test_import_creates_batch_leads_and_companies(self):
        product = Product.objects.create(name="Światłowód")
        upload = make_xlsx(
            [
                ["Anna Nowak", "600 100 200", "anna@example.com", "Alfa", "1234567890", "", "", "Olsztyn"],
                ["Piotr Zieliński", "+48 600-100-201", "", "Alfa", "1234567890", "", "", "Olsztyn"],
                ["Ewa Kot", 600100202, "", "", "", "", "", ""],
            ]
        )

        response = self.do_import(upload, name="Baza testowa", product=product.id)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["created"], 3)
        batch = LeadBatch.objects.get(name="Baza testowa")
        self.assertEqual(batch.source_file_name, "baza.xlsx")
        self.assertEqual(batch.created_by, self.manager)
        leads = Lead.objects.filter(batch=batch)
        self.assertEqual(leads.count(), 3)
        for lead in leads:
            self.assertIsNone(lead.owner)
            self.assertEqual(lead.status, Lead.Status.NEW)
            self.assertEqual(lead.product, product)
        self.assertEqual(Company.objects.filter(nip="1234567890").count(), 1)
        self.assertTrue(leads.filter(phone="600100202").exists())  # liczba z Excela

    # Duplikaty: numer już w systemie (inny zapis) i powtórzony w pliku
    def test_duplicates_are_skipped_and_reported(self):
        self.make_lead(phone="+48 600 100 200")
        upload = make_xlsx(
            [
                ["A", "600100200"],
                ["B", "600100203"],
                ["C", "600 100 203"],
            ],
            headers=["Imię i nazwisko", "Telefon"],
        )

        response = self.do_import(upload)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["created"], 1)
        rows = [item["row"] for item in response.data["duplicates"]]
        self.assertEqual(rows, [2, 4])

    # Błędne wiersze są raportowane z numerem, poprawne i tak się importują
    def test_invalid_rows_are_reported_with_row_numbers(self):
        upload = make_xlsx(
            [
                ["", "600100210"],
                ["X", "123"],
                ["Y", "600100211", "zly-mail"],
                ["Z", "600100212", "", "Firma", "12345"],
                ["Dobry", "600100213"],
            ]
        )

        response = self.do_import(upload)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["created"], 1)
        self.assertEqual([e["row"] for e in response.data["errors"]], [2, 3, 4, 5])
        self.assertEqual(response.data["total_rows"], 5)

    # Nagłówki rozpoznajemy też w innych zapisach
    def test_header_aliases_are_recognised(self):
        upload = make_xlsx([["Jan", "600100220"]], headers=["Osoba", "Nr telefonu"])

        response = self.do_import(upload)

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["created"], 1)

    # Brak wymaganych kolumn
    def test_missing_required_columns_returns_400(self):
        upload = make_xlsx([["Jan", "jan@example.com"]], headers=["Imię i nazwisko", "E-mail"])

        response = self.do_import(upload)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("telefon", response.data["detail"])

    # Plik, który nie jest Excelem, i zły format są odrzucane
    def test_invalid_files_are_rejected(self):
        wrong_ext = SimpleUploadedFile("baza.txt", b"abc", content_type="text/plain")
        broken = SimpleUploadedFile("baza.xlsx", b"to nie jest excel")

        self.assertEqual(
            self.do_import(wrong_ext).status_code, status.HTTP_400_BAD_REQUEST
        )
        self.assertEqual(self.do_import(broken).status_code, status.HTTP_400_BAD_REQUEST)

    # Gdy żaden wiersz nie jest poprawny, nie powstaje pusta paczka
    def test_no_valid_rows_creates_nothing(self):
        response = self.do_import(make_xlsx([["X", "123"], ["", "600100230"]]))

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(LeadBatch.objects.count(), 0)
        self.assertEqual(Lead.objects.count(), 0)

    # Nazwa bazy musi być unikalna
    def test_duplicate_batch_name_returns_400(self):
        LeadBatch.objects.create(name="Zajęta")

        response = self.do_import(make_xlsx([["Jan", "600100240"]]), name="Zajęta")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class ExcelExportTests(CrmTestBase):
    """Eksport leadów do Excela."""

    def read_sheet(self, response):
        workbook = load_workbook(io.BytesIO(response.content))
        return list(workbook.active.iter_rows(values_only=False))

    def setUp(self):
        super().setUp()
        company = Company.objects.create(name="Alfa", nip="1234567890")
        self.pool = self.make_lead(contact_name="Z Puli", phone="500000001")
        self.mine = self.make_lead(
            owner=self.agent1, contact_name="Mój", phone="500000002", company=company
        )
        self.lost = self.make_lead(
            owner=self.agent2,
            contact_name="Przegrany",
            phone="500000003",
            status=Lead.Status.LOST,
        )

    # Manager dostaje wszystkie leady w poprawnym pliku .xlsx
    def test_manager_exports_all_leads(self):
        response = self.api(self.manager).get("/api/leads/export/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        rows = self.read_sheet(response)
        self.assertEqual(rows[0][1].value, "Imię i nazwisko")
        self.assertEqual(len(rows), 4)  # nagłówek + 3 leady
        mine_row = next(r for r in rows if r[1].value == "Mój")
        self.assertEqual(mine_row[4].value, "Alfa")
        self.assertEqual(mine_row[10].value, "agent1")

    # Agent eksportuje tylko swoje leady
    def test_agent_exports_only_own_leads(self):
        response = self.api(self.agent1).get("/api/leads/export/")

        rows = self.read_sheet(response)
        self.assertEqual([r[1].value for r in rows[1:]], ["Mój"])

    # Eksport respektuje filtry z listy
    def test_export_respects_filters(self):
        response = self.api(self.manager).get("/api/leads/export/?status=lost")

        rows = self.read_sheet(response)
        self.assertEqual([r[1].value for r in rows[1:]], ["Przegrany"])

    # Dane zaczynające się od "=" trafiają do pliku jako tekst, nie formuła
    def test_formula_like_text_is_exported_as_text(self):
        self.make_lead(contact_name="=1+1", phone="500000004")

        response = self.api(self.manager).get("/api/leads/export/")

        cell = next(r[1] for r in self.read_sheet(response)[1:] if r[1].value == "=1+1")
        self.assertEqual(cell.data_type, "s")

    # Eksport wymaga logowania i roli
    def test_export_requires_login_and_role(self):
        self.assertEqual(
            self.api().get("/api/leads/export/").status_code,
            status.HTTP_401_UNAUTHORIZED,
        )
        self.assertEqual(
            self.api(self.no_role).get("/api/leads/export/").status_code,
            status.HTTP_403_FORBIDDEN,
        )


class TeamAccessFlagTests(CrmTestBase):
    """Flaga „może widzieć wyniki innych agentów” (AgentProfile) i jej zmiana przez managera."""

    def patch_flag(self, user, target, value):
        return self.api(user).patch(
            f"/api/agents/{target.pk}/", {"can_view_team_stats": value}, format="json"
        )

    # Bez profilu agent ma dostęp (domyślnie wszyscy widzą ranking)
    def test_agent_without_profile_can_view_team(self):
        self.assertTrue(can_view_team_stats(self.agent1))

    # Wyłączony profil blokuje agenta, ale nie managera
    def test_disabled_profile_blocks_agent_but_not_manager(self):
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        self.assertFalse(can_view_team_stats(self.agent1))
        self.assertTrue(can_view_team_stats(self.manager))

    # /api/me/ zwraca informację o dostępie do rankingu
    def test_me_returns_can_view_team(self):
        response = self.api(self.agent1).get("/api/me/")
        self.assertTrue(response.data["can_view_team"])
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        response = self.api(self.agent1).get("/api/me/")
        self.assertFalse(response.data["can_view_team"])

    # Lista agentów (manager) pokazuje flagę
    def test_agents_list_shows_flag(self):
        AgentProfile.objects.create(user=self.agent2, can_view_team_stats=False)
        rows = {r["username"]: r for r in self.api(self.manager).get("/api/agents/").data}
        self.assertTrue(rows["agent1"]["can_view_team_stats"])
        self.assertFalse(rows["agent2"]["can_view_team_stats"])

    # Manager wyłącza i włącza podgląd agentowi
    def test_manager_can_toggle_flag(self):
        response = self.patch_flag(self.manager, self.agent1, False)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(can_view_team_stats(self.agent1))
        self.patch_flag(self.manager, self.agent1, True)
        self.assertTrue(can_view_team_stats(self.agent1))
        self.assertEqual(AgentProfile.objects.filter(user=self.agent1).count(), 1)

    # Agent nie może zmienić flagi ani sobie, ani innym (403)
    def test_agent_cannot_change_flag(self):
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        own = self.patch_flag(self.agent1, self.agent1, True)
        other = self.patch_flag(self.agent1, self.agent2, False)
        self.assertEqual(own.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(other.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(can_view_team_stats(self.agent1))

    # Bez tokena 401, błędne dane 400, managerowi nie da się wyłączyć podglądu
    def test_patch_flag_validation(self):
        self.assertEqual(
            self.api().patch(f"/api/agents/{self.agent1.pk}/", {}, format="json").status_code,
            status.HTTP_401_UNAUTHORIZED,
        )
        empty = self.api(self.manager).patch(
            f"/api/agents/{self.agent1.pk}/", {}, format="json"
        )
        self.assertEqual(empty.status_code, status.HTTP_400_BAD_REQUEST)
        other_manager = User.objects.create_user("manager2", password="haslo12345")
        other_manager.groups.add(Group.objects.get(name=MANAGER))
        self.assertEqual(
            self.patch_flag(self.manager, other_manager, False).status_code,
            status.HTTP_400_BAD_REQUEST,
        )


class LeaderboardTests(CrmTestBase):
    """Ranking agentów, cudze statystyki i blokada 403 ustawiana przez managera."""

    def add_calls(self, agent, outcome, count, created_at=None):
        lead = self.make_lead(owner=agent)
        for _ in range(count):
            activity = Activity.objects.create(lead=lead, agent=agent, outcome=outcome)
            if created_at:
                Activity.objects.filter(pk=activity.pk).update(created_at=created_at)

    def ranking(self, user, query=""):
        return self.api(user).get(f"/api/leaderboard/{query}")

    # Ranking jest posortowany wg sprzedanych, a remis rozstrzyga liczba rozmów
    def test_ranking_sorted_by_sold_then_calls(self):
        self.add_calls(self.agent1, Activity.Outcome.SOLD, 1)
        self.add_calls(self.agent2, Activity.Outcome.SOLD, 1)
        self.add_calls(self.agent2, Activity.Outcome.NO_ANSWER, 3)
        response = self.ranking(self.agent1, "?days=1")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = response.data["results"]
        self.assertEqual([r["username"] for r in rows], ["agent2", "agent1"])
        self.assertEqual([r["rank"] for r in rows], [1, 2])
        self.assertEqual(rows[0]["sold"], 1)
        self.assertEqual(rows[0]["calls"], 4)

    # Ranking zwraca zdjęcie profilowe agenta (albo pusty tekst, gdy go nie ma)
    def test_ranking_rows_include_avatar(self):
        from .models import AgentProfile

        png = "data:image/png;base64,AAAA"
        AgentProfile.objects.create(user=self.agent1, avatar=png)
        rows = {r["username"]: r for r in self.ranking(self.manager, "?days=1").data["results"]}
        self.assertEqual(rows["agent1"]["avatar"], png)
        self.assertEqual(rows["agent2"]["avatar"], "")

    # Taki sam wynik i liczba rozmów dają to samo miejsce
    def test_equal_scores_share_rank(self):
        self.add_calls(self.agent1, Activity.Outcome.SOLD, 1)
        self.add_calls(self.agent2, Activity.Outcome.SOLD, 1)
        rows = self.ranking(self.manager, "?days=1").data["results"]
        self.assertEqual([r["rank"] for r in rows], [1, 1])

    # Własny wiersz jest oznaczony, a managera nie ma w rankingu
    def test_own_row_marked_and_manager_not_listed(self):
        rows = self.ranking(self.agent1).data["results"]
        self.assertEqual({r["username"] for r in rows}, {"agent1", "agent2"})
        self.assertEqual([r["username"] for r in rows if r["is_me"]], ["agent1"])

    # Ranking można układać wg rozmów i zainteresowanych
    def test_sort_by_calls_and_interested(self):
        self.add_calls(self.agent1, Activity.Outcome.INTERESTED, 1)
        self.add_calls(self.agent2, Activity.Outcome.NO_ANSWER, 5)
        by_calls = self.ranking(self.manager, "?days=1&sort=calls").data["results"]
        by_interest = self.ranking(self.manager, "?days=1&sort=interested").data["results"]
        self.assertEqual(by_calls[0]["username"], "agent2")
        self.assertEqual(by_interest[0]["username"], "agent1")
        bad = self.ranking(self.manager, "?sort=pieniadze")
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    # Okres „dziś” liczy tylko dzisiejsze rozmowy
    def test_today_period_counts_only_today(self):
        old = timezone.now() - timedelta(days=3)
        self.add_calls(self.agent1, Activity.Outcome.SOLD, 2, created_at=old)
        self.add_calls(self.agent2, Activity.Outcome.SOLD, 1)
        today = self.ranking(self.manager, "?days=1").data["results"]
        week = self.ranking(self.manager, "?days=7").data["results"]
        self.assertEqual({r["username"]: r["sold"] for r in today}, {"agent1": 0, "agent2": 1})
        self.assertEqual({r["username"]: r["sold"] for r in week}, {"agent1": 2, "agent2": 1})

    # Zakres dat i błędne okresy działają jak w /api/stats/
    def test_ranking_period_validation(self):
        today = timezone.localdate().isoformat()
        ok = self.ranking(self.manager, f"?date_from={today}&date_to={today}")
        self.assertEqual(ok.status_code, status.HTTP_200_OK)
        bad = self.ranking(self.manager, "?days=0")
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    # Ranking nie zawiera danych klientów ani leadów
    def test_ranking_has_no_client_data(self):
        self.add_calls(self.agent2, Activity.Outcome.SOLD, 1, created_at=None)
        text = str(self.ranking(self.agent1, "?days=1").data)
        self.assertNotIn("500100200", text)
        self.assertNotIn("Kowalski", text)

    # Niezalogowany 401, użytkownik bez roli 403
    def test_ranking_requires_login_and_role(self):
        self.assertEqual(self.ranking(None).status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.ranking(self.no_role).status_code, status.HTTP_403_FORBIDDEN)

    # Agent z wyłączonym podglądem dostaje 403 z czytelnym komunikatem
    def test_blocked_agent_gets_403_on_ranking(self):
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        response = self.ranking(self.agent1)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn("wyłączony przez managera", response.data["detail"])
        # inny agent i manager nadal widzą ranking
        self.assertEqual(self.ranking(self.agent2).status_code, status.HTTP_200_OK)
        self.assertEqual(self.ranking(self.manager).status_code, status.HTTP_200_OK)

    # Zablokowany agent nadal jest widoczny w rankingu u innych
    def test_blocked_agent_still_listed_for_others(self):
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        rows = self.ranking(self.agent2).data["results"]
        self.assertIn("agent1", [r["username"] for r in rows])

    # Agent z dostępem widzi liczby innego agenta, ale nie stan jego leadów
    def test_agent_sees_peer_numbers_without_lead_data(self):
        self.add_calls(self.agent2, Activity.Outcome.SOLD, 2)
        response = self.api(self.agent1).get(f"/api/stats/?days=1&agent={self.agent2.pk}")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["outcomes"]["sold"], 2)
        self.assertEqual(response.data["agents"][0]["username"], "agent2")
        self.assertEqual(response.data["total_leads"], 0)
        self.assertEqual(sum(response.data["leads_by_status"].values()), 0)
        self.assertEqual(sum(response.data["scheduled"].values()), 0)
        self.assertIsNone(response.data["pool_leads"])

    # Agent bez dostępu dostaje 403 na cudze statystyki, ale własne działają
    def test_blocked_agent_gets_403_on_peer_stats(self):
        AgentProfile.objects.create(user=self.agent1, can_view_team_stats=False)
        peer = self.api(self.agent1).get(f"/api/stats/?agent={self.agent2.pk}")
        own = self.api(self.agent1).get(f"/api/stats/?agent={self.agent1.pk}")
        plain = self.api(self.agent1).get("/api/stats/")
        self.assertEqual(peer.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(own.status_code, status.HTTP_200_OK)
        self.assertEqual(plain.status_code, status.HTTP_200_OK)

    # Manager widzi pełne statystyki wskazanego agenta; zły parametr to 400/404
    def test_manager_stats_for_agent_and_bad_param(self):
        self.make_lead(owner=self.agent2)
        response = self.api(self.manager).get(f"/api/stats/?agent={self.agent2.pk}")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_leads"], 1)
        bad = self.api(self.manager).get("/api/stats/?agent=abc")
        missing = self.api(self.manager).get("/api/stats/?agent=99999")
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(missing.status_code, status.HTTP_404_NOT_FOUND)

    # Po włączeniu podglądu przez managera agent znowu widzi ranking
    def test_manager_unblocking_restores_access(self):
        self.api(self.manager).patch(
            f"/api/agents/{self.agent1.pk}/", {"can_view_team_stats": False}, format="json"
        )
        self.assertEqual(self.ranking(self.agent1).status_code, status.HTTP_403_FORBIDDEN)
        self.api(self.manager).patch(
            f"/api/agents/{self.agent1.pk}/", {"can_view_team_stats": True}, format="json"
        )
        self.assertEqual(self.ranking(self.agent1).status_code, status.HTTP_200_OK)


class BulkTeamAccessTests(CrmTestBase):
    """Jeden przycisk managera: włączenie lub wyłączenie rankingu wszystkim agentom naraz."""

    def bulk(self, user, value):
        return self.api(user).post(
            "/api/agents/team-access/", {"can_view_team_stats": value}, format="json"
        )

    def test_manager_disables_ranking_for_all_agents(self):
        response = self.bulk(self.manager, False)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data["can_view_team_stats"])
        self.assertEqual(response.data["updated"], 2)
        self.assertFalse(can_view_team_stats(self.agent1))
        self.assertFalse(can_view_team_stats(self.agent2))
        self.assertEqual(
            self.api(self.agent1).get("/api/leaderboard/").status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(self.api(self.manager).get("/api/leaderboard/").status_code, 200)

    def test_manager_enables_ranking_again_for_all(self):
        self.bulk(self.manager, False)
        self.bulk(self.manager, True)
        self.assertTrue(can_view_team_stats(self.agent1))
        self.assertTrue(can_view_team_stats(self.agent2))
        self.assertEqual(self.api(self.agent2).get("/api/leaderboard/").status_code, 200)

    def test_agent_cannot_use_bulk_switch(self):
        self.assertEqual(self.bulk(self.agent1, False).status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(can_view_team_stats(self.agent2))

    def test_anonymous_gets_401(self):
        self.assertEqual(self.bulk(None, False).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_bulk_requires_boolean_value(self):
        response = self.api(self.manager).post("/api/agents/team-access/", {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class LeadLookupTests(CrmTestBase):
    """Wyszukiwarka leadów po numerze telefonu, NIP i nazwisku."""

    def setUp(self):
        super().setUp()
        self.company = Company.objects.create(name="Alfa Sp. z o.o.", nip="5250000003")
        self.mine = self.make_lead(
            owner=self.agent1, contact_name="Anna Nowak", phone="500 100 200", company=self.company
        )
        self.other = self.make_lead(owner=self.agent2, contact_name="Cudzy Lead", phone="600700800")

    def lookup(self, user, q):
        return self.api(user).get("/api/leads/lookup/", {"q": q})

    def test_finds_lead_by_phone_ignoring_spaces_and_country_prefix(self):
        for q in ("500100200", "500 100 200", "+48 500-100-200", "0048500100200"):
            response = self.lookup(self.agent1, q)
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertEqual([r["id"] for r in response.data], [self.mine.id], q)
            self.assertEqual(response.data[0]["match"], "exact")
            self.assertEqual(response.data[0]["matched_on"], "phone")

    def test_finds_lead_by_company_nip(self):
        response = self.lookup(self.agent1, "5250000003")
        self.assertEqual([r["id"] for r in response.data], [self.mine.id])
        self.assertEqual(response.data[0]["match"], "exact")
        self.assertEqual(response.data[0]["matched_on"], "nip")

    def test_partial_digits_and_name_search(self):
        self.assertEqual([r["id"] for r in self.lookup(self.agent1, "100200").data], [self.mine.id])
        self.assertEqual(self.lookup(self.agent1, "100200").data[0]["match"], "partial")
        response = self.lookup(self.agent1, "nowak")
        self.assertEqual([r["id"] for r in response.data], [self.mine.id])
        self.assertEqual(response.data[0]["matched_on"], "name")

    def test_agent_does_not_find_someone_elses_lead(self):
        self.assertEqual(self.lookup(self.agent1, "600700800").data, [])

    def test_manager_finds_any_lead(self):
        self.assertEqual([r["id"] for r in self.lookup(self.manager, "600700800").data], [self.other.id])

    def test_too_short_query_returns_nothing(self):
        self.assertEqual(self.lookup(self.agent1, "5").data, [])
        self.assertEqual(self.lookup(self.agent1, "50").data, [])

    def test_requires_role(self):
        self.assertEqual(self.lookup(self.no_role, "500").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.lookup(None, "500").status_code, status.HTTP_401_UNAUTHORIZED)



class ResetLeadsCommandTests(CrmTestBase):
    """Komenda reset_leads: czyści leady do testów, ale tylko w trybie DEBUG."""

    def run_command(self, *args):
        from io import StringIO

        from django.core.management import call_command

        out = StringIO()
        call_command("reset_leads", *args, stdout=out)
        return out.getvalue()

    def test_refuses_when_debug_is_off(self):
        from django.core.management.base import CommandError

        self.make_lead(owner=self.agent1)
        with override_settings(DEBUG=False):
            with self.assertRaises(CommandError):
                self.run_command("--yes")
        self.assertEqual(Lead.objects.count(), 1)

    @override_settings(DEBUG=True)
    def test_empty_reset_removes_leads_but_keeps_users(self):
        lead = self.make_lead(owner=self.agent1)
        Activity.objects.create(lead=lead, agent=self.agent1, outcome="no_answer")
        self.run_command("--yes", "--empty")
        self.assertEqual(Lead.objects.count(), 0)
        self.assertEqual(Activity.objects.count(), 0)
        self.assertTrue(User.objects.filter(username="agent1").exists())

    @override_settings(DEBUG=True)
    def test_reset_loads_demo_data(self):
        self.make_lead(owner=self.agent1, phone="500111222")
        self.run_command("--yes")
        self.assertFalse(Lead.objects.filter(phone="500111222").exists())
        self.assertGreater(Lead.objects.count(), 5)

    @override_settings(DEBUG=True)
    def test_wrong_confirmation_cancels(self):
        from unittest import mock

        self.make_lead(owner=self.agent1)
        with mock.patch("builtins.input", return_value="nie"):
            out = self.run_command()
        self.assertIn("Anulowano", out)
        self.assertEqual(Lead.objects.count(), 1)
